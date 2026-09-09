"""Parse independent sheets by labels and headers, never by fixed cell positions."""

from dataclasses import dataclass
from datetime import date
from pathlib import Path
import re
from zipfile import BadZipFile, ZipFile

import openpyxl
import config
from .models import Invoice, InvoiceControl, Project, RegisterTask, ReportData, Task, WorkbookError
from .normalization import (
    MISSING,
    clean_text,
    explicit_overdue,
    explicit_due_date,
    is_paid,
    key,
    money_mentions,
    normalize_currency,
    normalize_date,
    normalize_name,
    normalize_percentage,
    normalize_project_name,
    normalize_status,
    reporting_period,
    safe_number,
    section_category,
    stale_context,
)

COLUMN_ALIASES = {
    "project": ["project", "project name", "client project"],
    "task_id": ["task id", "taskid", "id", "task number"],
    "title": ["task title", "title", "task / deliverable", "task", "deliverable", "description"],
    "status": ["status", "task status", "current status", "payment status"],
    "priority": ["priority", "urgency"],
    "assignee": ["assignee", "assigned to", "owner", "team member"],
    "created": ["created", "created date", "date created", "created on"],
    "due": ["due", "due date", "task due date", "deadline"],
    "updated": ["updated", "last updated", "updated at", "modified date"],
    "notes": [
        "notes / flag",
        "notes",
        "flag",
        "flags",
        "comments",
        "remarks",
        "ongoing awaiting tasks",
    ],
    "overdue_flag": ["overdue flag", "is overdue"],
    "start": ["project start date", "start date", "start"],
    "end": ["project end date", "end date", "end", "target date"],
    "overall": [
        "completion %",
        "overall completion",
        "overall completion %",
        "project completion %",
        "percent complete",
    ],
    "total": ["total tasks (this week)", "total tasks", "tasks this week", "weekly tasks"],
    "done": ["done", "completed", "done this week", "completed tasks"],
    "in_progress": ["in progress", "started"],
    "in_review": ["in review", "review"],
    "hold_todo": ["on hold / to do", "on hold / todo", "on hold and to do"],
    "on_hold": ["on hold", "held"],
    "todo": ["to do", "todo", "not started"],
    "overdue": ["overdue", "overdue tasks"],
    "client": ["client", "customer", "client name"],
    "invoice_id": ["invoice #", "invoice number", "invoice no", "invoice id", "invoice"],
    "amount": ["amount", "invoice amount", "value", "deal value", "amount due"],
    "currency": ["currency", "ccy"],
    "phase": ["phase", "project phase"],
    "sprint": ["sprint", "sprints", "sprint number"],
    "hours": ["est. hours", "estimated hours", "est hours", "effort hours", "hours"],
    "month": ["month", "stage month"],
    "opportunity": ["opportunity", "project / opportunity", "opportunity name", "proposal"],
    "proposal_date": ["proposal date", "date proposed", "date sent", "sent date"],
    "stage": ["stage", "sales stage", "lead status"],
    "next_action": ["next action", "action required", "follow up"],
    "advance_status": ["advance status", "advance payment status"],
    "section": ["category", "section", "invoice category", "bucket"],
}
ALIASES = {key(alias): name for name, aliases in COLUMN_ALIASES.items() for alias in aliases}
COUNTS = ("total", "done", "in_progress", "in_review", "hold_todo", "on_hold", "todo", "overdue")


@dataclass
class Sheet:
    title: str
    rows: list

    def source(self, row: int) -> str:
        return f"{self.title}, row {row}"


def columns(row) -> dict[str, int]:
    return {ALIASES[key(cell.value)]: i for i, cell in enumerate(row) if key(cell.value) in ALIASES}


def get(row, mapping, field, default=None):
    i = mapping.get(field)
    return row[i].value if i is not None and i < len(row) else default


def find_header(sheet: Sheet, required: set[str]):
    for i, row in enumerate(sheet.rows):
        mapping = columns(row)
        if required <= mapping.keys():
            return i, mapping
    return None


def find_sheet(sheets, preferred, required, excluded=()):
    candidates = [s for s in sheets if s.title not in excluded]
    candidates.sort(key=lambda s: (not any(key(p) in key(s.title) for p in preferred),))
    return next((s for s in candidates if find_header(s, required)), None)


def parse_date(value, data: ReportData, source: str, label: str, *, past=False):
    parsed = normalize_date(value, data.report_end_date, past=past)
    if clean_text(value).casefold() not in MISSING and parsed is None:
        data.warn(
            f"Invalid or unavailable {label} at {source}: {clean_text(value)}. No date inferred."
        )
    return parsed


def canonical_project(name, data):
    normalized = normalize_project_name(name)
    existing = next((p.name for p in data.projects if key(p.name) == key(normalized)), None)
    if existing:
        return existing
    data.projects.append(Project(name=normalized))
    return normalized


def parse_metadata(sheets: list[Sheet], data: ReportData) -> None:
    candidates = []
    separate_start = separate_end = None
    for sheet in sheets:
        for row in sheet.rows:
            values = [c.value for c in row if c.value is not None]
            joined = " | ".join(clean_text(v) for v in values)
            for i, value in enumerate(values):
                raw = clean_text(value)
                if re.search(
                    r"reporting period|report period|detailed task log|week of", raw, re.I
                ):
                    start, end = reporting_period(" ".join(clean_text(v) for v in values[i:]))
                    if start and end:
                        candidates.append(
                            (0 if "reporting period" in raw.casefold() else 1, start, end)
                        )
                if key(raw) in ("reportstartdate", "reportingstartdate") and i + 1 < len(values):
                    separate_start = normalize_date(values[i + 1])
                if key(raw) in ("reportenddate", "reportingenddate") and i + 1 < len(values):
                    separate_end = normalize_date(values[i + 1])
            prepared = re.search(r"prepared\s*by\s*:?\s*(.+?)(?:\||$)", joined, re.I)
            if prepared and not data.prepared_by:
                data.prepared_by = prepared[1].strip(" ,|:")
            # Accommodate a label and value in separate cells.
            for i, value in enumerate(values[:-1]):
                if key(value) == "preparedby":
                    data.prepared_by = clean_text(values[i + 1]).strip(" ,")
                if key(value) in ("preparedbyrole", "role", "designation"):
                    data.prepared_role = clean_text(values[i + 1])
    if separate_start and separate_end and separate_start <= separate_end:
        candidates.insert(0, (-1, separate_start, separate_end))
    if candidates:
        candidates.sort(key=lambda item: item[0])
        _, data.report_start_date, data.report_end_date = candidates[0]
        if len({(a, b) for _, a, b in candidates}) > 1:
            data.warn(
                "Conflicting reporting periods in workbook headings. The explicit Reporting Period label takes precedence."
            )
    else:
        data.warn(
            "Missing reporting period. Date-based overdue, lead aging and timeline checks are unavailable."
        )


def parse_executive_summary(sheet: Sheet | None, data: ReportData) -> None:
    if sheet is None:
        data.warn(
            "Executive Summary unavailable. Overall project completion and project dates may be missing."
        )
        return
    located = find_header(sheet, {"project"})
    if not located:
        return
    header, mapping = located
    data.sources["executive_summary"] = sheet.title
    for n, row in enumerate(sheet.rows[header + 1 :], header + 2):
        if {"project", "total"} <= columns(row).keys():
            mapping = columns(row)
            continue
        name = clean_text(get(row, mapping, "project"))
        if not name:
            continue
        if key(name) in ("total", "grandtotal", "portfoliototal"):
            data.reported_totals = {k: safe_number(get(row, mapping, k)) for k in COUNTS}
            break
        if re.match(r"status legend|note\s*:|legend", name, re.I):
            break
        if len([c for c in row if clean_text(c.value)]) == 1 and not any(
            k in mapping for k in ("total", "overall", "start")
        ):
            continue
        source = sheet.source(n)
        overall_cell = row[mapping["overall"]] if "overall" in mapping else None
        overall = (
            normalize_percentage(overall_cell.value, overall_cell.number_format)
            if overall_cell
            else None
        )
        if overall_cell and clean_text(overall_cell.value) and overall is None:
            data.warn(f"Invalid overall completion for {name} at {source}.")
        notes = clean_text(get(row, mapping, "notes"))
        project = Project(
            normalize_project_name(name),
            overall,
            parse_date(get(row, mapping, "start"), data, source, "project start date"),
            parse_date(get(row, mapping, "end"), data, source, "project end date"),
            notes,
            {k: safe_number(get(row, mapping, k)) for k in COUNTS},
            stale_context(notes),
            source,
        )
        existing = next((p for p in data.projects if key(p.name) == key(project.name)), None)
        if existing:
            data.warn(f"Duplicate project summary: {project.name}. Kept the last row ({source}).")
            data.projects.remove(existing)
        data.projects.append(project)


def task_from_row(row, mapping, project, data, source):
    title = clean_text(get(row, mapping, "title"))
    task_id = clean_text(get(row, mapping, "task_id"))
    status_raw = clean_text(get(row, mapping, "status"))
    if (
        not title
        and not task_id
        or key(task_id) in ("total", "grandtotal")
        or key(title) in ("total", "grandtotal")
    ):
        return None
    if title and not task_id and not status_raw and not clean_text(get(row, mapping, "assignee")):
        if re.match(r"note\s*:|status legend|legend", title, re.I):
            return None
    project = canonical_project(project, data)
    if project == "Unspecified project":
        data.warn(f"Missing project at {source}; grouped as Unspecified project.")
    status = normalize_status(status_raw)
    if status == "Unknown":
        data.warn(
            f"Unrecognized task status '{status_raw or '(blank)'}' for {task_id or title} ({source})."
        )
    assignee = normalize_name(get(row, mapping, "assignee"))
    # Merge casing/spacing variants without guessing that different names are one person.
    assignee = next((t.assignee for t in data.tasks if key(t.assignee) == key(assignee)), assignee)
    notes = clean_text(get(row, mapping, "notes"))
    due = parse_date(get(row, mapping, "due"), data, source, "task due date")
    overdue_cell = get(row, mapping, "overdue_flag", get(row, mapping, "overdue"))
    flag = explicit_overdue(notes) or key(overdue_cell) in ("yes", "true", "1", "overdue")
    overdue = flag or bool(
        due and data.report_end_date and due < data.report_end_date and status != "Done"
    )
    stale = stale_context(notes) or next(p.stale for p in data.projects if p.name == project)
    return Task(
        project,
        task_id,
        title or "Not available",
        status,
        clean_text(get(row, mapping, "priority")),
        assignee,
        parse_date(get(row, mapping, "created"), data, source, "created date", past=True),
        due,
        parse_date(get(row, mapping, "updated"), data, source, "updated date", past=True),
        notes,
        overdue,
        stale,
        status_raw,
        source,
    )


def deduplicate_tasks(tasks, data):
    kept = {}
    for index, task in enumerate(tasks):
        identity = key(task.task_id) if task.task_id else f"anonymous:{index}"
        if identity in kept:
            previous = kept[identity]

            def rank(t):
                fields = (
                    t.title != "Not available",
                    t.status != "Unknown",
                    t.assignee != "Unassigned",
                    t.created_date,
                    t.due_date,
                    t.notes,
                    t.project != "Unspecified project",
                )
                return t.updated_date or date.min, sum(bool(v) for v in fields)

            winner = task if rank(task) >= rank(previous) else previous
            data.warn(
                f"Duplicate Task ID {task.task_id}: counted once. Kept {winner.source}; latest update, then most complete record, then last row wins. Review conflicting copies."
            )
            kept[identity] = winner
        else:
            kept[identity] = task
    return list(kept.values())


def parse_task_log(sheet: Sheet | None, data: ReportData) -> list[Task]:
    if sheet is None:
        return []
    located = find_header(sheet, {"project", "title", "status"})
    if not located:
        return []
    data.task_source_available = True
    data.sources["weekly_tasks"] = sheet.title
    header, mapping = located
    tasks = []
    for n, row in enumerate(sheet.rows[header + 1 :], header + 2):
        if {"title", "status"} <= columns(row).keys():
            mapping = columns(row)
            continue
        task = task_from_row(row, mapping, get(row, mapping, "project"), data, sheet.source(n))
        if task:
            tasks.append(task)
    return deduplicate_tasks(tasks, data)


def parse_project_breakdown(sheet: Sheet | None, data: ReportData) -> list[Task]:
    if sheet is None:
        return []
    tasks, mapping, current_project = [], None, ""
    for n, row in enumerate(sheet.rows, 1):
        text = " ".join(clean_text(c.value) for c in row if clean_text(c.value))
        if re.search(r"\|\s*\d+\s*/\s*\d+\s*tasks?", text, re.I):
            current_project = text.split("|")[0].strip()
            current_project = re.sub(
                r"\s*\((?:no data this week|\d+% complete)\)|\s*[★*]\s*new\b",
                "",
                current_project,
                flags=re.I,
            ).strip()
            continue
        if {"title", "status"} <= columns(row).keys():
            mapping = columns(row)
            continue
        if mapping and current_project:
            task = task_from_row(row, mapping, current_project, data, sheet.source(n))
            if task:
                tasks.append(task)
    if tasks:
        data.sources["project_breakdown"] = sheet.title
    return deduplicate_tasks(tasks, data)


def parse_invoice_summary(sheet: Sheet | None, data: ReportData) -> None:
    if sheet is None:
        return
    data.sources["invoices"] = sheet.title
    data.finance_stale = bool(
        re.search(
            r"carried|no new payment|no fresh",
            sheet.title + " ".join(clean_text(c.value) for row in sheet.rows for c in row),
            re.I,
        )
    )
    if data.finance_stale:
        data.warn(
            "Finance data appears carried forward. Calculated classifications use the report end date and recorded line statuses; confirm payment freshness."
        )
    mapping, section, category, pending_sources = None, "Invoice records", "Unclassified", []
    exact_records = set()
    for n, row in enumerate(sheet.rows, 1):
        found = columns(row)
        if {"client", "amount", "status"} <= found.keys():
            mapping = found
            continue
        values = [clean_text(c.value) for c in row if clean_text(c.value)]
        if not values:
            continue
        first = values[0]
        if re.match(r"note\s*:", first, re.I):
            continue
        if len(values) == 1 and section_category(first) != "Unclassified":
            section, category, pending_sources = first, section_category(first), []
            continue
        if not mapping:
            continue
        source = sheet.source(n)
        client = clean_text(get(row, mapping, "client"))
        if re.match(r"sub\s*total|grand\s*total|^total\b", client, re.I):
            amount = safe_number(get(row, mapping, "amount"))
            currency = normalize_currency(
                get(row, mapping, "currency"), get(row, mapping, "amount")
            )
            notes = (
                clean_text(get(row, mapping, "status"))
                + " "
                + clean_text(get(row, mapping, "notes"))
            )
            data.invoice_controls.append(
                InvoiceControl(
                    client,
                    {currency: amount} if amount is not None else {},
                    list(pending_sources),
                    notes.strip(),
                    source,
                )
            )
            pending_sources = []
            continue
        raw_amount = get(row, mapping, "amount")
        invoice_id = clean_text(get(row, mapping, "invoice_id"))
        status = clean_text(get(row, mapping, "status"))
        if not client or not invoice_id and raw_amount is None:
            continue
        amount = safe_number(raw_amount)
        currency = normalize_currency(get(row, mapping, "currency"), raw_amount)
        if amount is None:
            data.warn(
                f"Missing or invalid invoice amount for {client} at {source}. Omitted from monetary totals."
            )
        if currency == "Unspecified currency":
            data.warn(
                f"Missing currency for {client} at {source}. Amount is kept in an unspecified-currency group."
            )
        notes = clean_text(get(row, mapping, "notes"))
        paid = is_paid(status)
        installment = (
            "Advance"
            if re.search(r"\badvance\b", status, re.I)
            else "Balance" if re.search(r"\bbalance\b", status, re.I) else "Invoice"
        )
        conditional = (
            installment == "Balance"
            and bool(re.search(r"on (?:project )?completion|on delivery|on handover", status, re.I))
            and not paid
        )
        cat_value = clean_text(get(row, mapping, "section"))
        record = Invoice(
            client,
            invoice_id,
            amount,
            currency,
            status,
            cat_value or section,
            section_category(cat_value) if cat_value else category,
            parse_date(get(row, mapping, "due"), data, source, "invoice due date"),
            notes,
            installment,
            paid,
            installment == "Advance" and not paid,
            conditional,
            data.finance_stale,
            source,
        )
        if record.due_date is None:
            record.due_date = explicit_due_date(status, data.report_end_date)
        identity = (
            key(client),
            key(invoice_id),
            installment,
            amount,
            currency,
            record.due_date,
            status,
        )
        if invoice_id and identity in exact_records:
            data.warn(
                f"Exact duplicate invoice line at {source}. Counted once; distinct clients and installments remain separate."
            )
            continue
        exact_records.add(identity)
        data.invoices.append(record)
        pending_sources.append(source)
        if paid and record.section_category in (
            "Overdue / Collections",
            "Due Soon",
            "Future Unpaid",
        ):
            data.warn(
                f"Finance reconciliation required: {client}, invoice {invoice_id} is Paid inside {record.section_category}. Excluded from outstanding ({source})."
            )
        status_date = (
            explicit_due_date(status, data.report_end_date)
            if re.search(r"\bdue\s+\d", status, re.I)
            else None
        )
        if status_date and record.due_date and status_date != record.due_date:
            data.warn(
                f"Invoice due-date conflict for {client} / {invoice_id}: Due Date column {record.due_date} versus status {status_date}. Used Due Date column."
            )
        elif re.search(r"\bdue\s+\d", status, re.I) and not status_date:
            data.warn(
                f"Malformed date in invoice status for {client} / {invoice_id}: {status}. Used the Due Date column where available."
            )


def parse_task_register(sheet: Sheet | None, data: ReportData) -> None:
    if sheet is None:
        return
    located = find_header(sheet, {"title", "status", "phase"})
    if not located:
        return
    header, mapping = located
    heading = " ".join(clean_text(c.value) for row in sheet.rows[:header] for c in row)
    matches = [p.name for p in data.projects if key(p.name) and key(p.name) in key(heading)]
    selected = config.REGISTER_PROJECTS.get(sheet.title, "")
    if not selected and len(matches) == 1:
        selected = matches[0]
    if not selected and "project" not in mapping:
        data.warn(
            f"{sheet.title}: register project could not be identified. Add a Project column or project name in the heading, or set REGISTER_PROJECTS once. Unlinked register rows are not assigned to a spotlight."
        )
    data.sources["task_register"] = sheet.title
    seen = set()
    for n, row in enumerate(sheet.rows[header + 1 :], header + 2):
        if {"title", "status", "phase"} <= columns(row).keys():
            mapping = columns(row)
            continue
        first = next((clean_text(c.value) for c in row if clean_text(c.value)), "")
        if key(first) in ("total", "grandtotal"):
            data.register_totals.append(
                {
                    "project": selected,
                    "hours": safe_number(get(row, mapping, "hours")),
                    "source": sheet.source(n),
                }
            )
            continue
        title = clean_text(get(row, mapping, "title"))
        if not title:
            continue
        project_value = get(row, mapping, "project") or selected
        project = canonical_project(project_value, data) if project_value else "Unlinked register"
        source = sheet.source(n)
        task_id = clean_text(get(row, mapping, "task_id"))
        if not task_id:
            # The register may use a numbered '#' column; it is an identifier only here.
            number_col = next(
                (i for i, c in enumerate(sheet.rows[header]) if clean_text(c.value) == "#"), None
            )
            task_id = clean_text(row[number_col].value) if number_col is not None else ""
        identity = (key(project), key(task_id)) if task_id else None
        if identity and identity in seen:
            data.warn(
                f"Duplicate register Task ID {task_id} in {project}. First record retained; review {source}."
            )
            continue
        if identity:
            seen.add(identity)
        raw = clean_text(get(row, mapping, "status"))
        status = normalize_status(raw)
        hours = safe_number(get(row, mapping, "hours"))
        if hours is not None and hours < 0:
            hours = None
        if hours is None:
            data.warn(
                f"Missing or invalid estimated hours for register task {task_id or title} ({source}). Hours totals cover populated nonnegative estimates only."
            )
        if status == "Unknown":
            data.warn(
                f"Unrecognized register status '{raw or '(blank)'}' for {task_id or title} ({source}). Hours kept under Unknown, not assumed completed or active."
            )
        data.register.append(
            RegisterTask(
                project,
                task_id,
                title,
                clean_text(get(row, mapping, "phase")) or "Not available",
                clean_text(get(row, mapping, "sprint")),
                clean_text(get(row, mapping, "month")),
                status,
                hours,
                normalize_name(get(row, mapping, "assignee")),
                clean_text(get(row, mapping, "notes")),
                raw,
                source,
            )
        )


def parse_sales_sheets(sheets, data):
    """Optional explicit pipeline tables; commercial invoices and proposal notes are handled by analytics."""
    records = []
    for sheet in sheets:
        located = find_header(sheet, {"client", "opportunity"})
        if not located:
            continue
        header, mapping = located
        for n, row in enumerate(sheet.rows[header + 1 :], header + 2):
            if {"client", "opportunity"} <= columns(row).keys():
                continue
            client, opportunity = clean_text(get(row, mapping, "client")), clean_text(
                get(row, mapping, "opportunity")
            )
            if not client and not opportunity or key(client) in ("total", "subtotal"):
                continue
            amount = safe_number(get(row, mapping, "amount"))
            currency = normalize_currency(
                get(row, mapping, "currency"), get(row, mapping, "amount")
            )
            records.append(
                {
                    "client": client or "Not available",
                    "opportunity": opportunity or "Not available",
                    "value": {currency: amount} if amount is not None else {},
                    "stage": clean_text(get(row, mapping, "stage") or get(row, mapping, "status")),
                    "proposal_date": parse_date(
                        get(row, mapping, "proposal_date"),
                        data,
                        sheet.source(n),
                        "proposal date",
                        past=True,
                    ),
                    "next_action": clean_text(get(row, mapping, "next_action")),
                    "advance_status": clean_text(get(row, mapping, "advance_status")),
                    "invoice_id": clean_text(get(row, mapping, "invoice_id")),
                    "source": sheet.source(n),
                    "notes": clean_text(get(row, mapping, "notes")),
                }
            )
    return records


def parse_workbook(path: str | Path) -> ReportData:
    path = Path(path)
    if path.suffix.casefold() != ".xlsx":
        raise WorkbookError("Please select an .xlsx Excel workbook.")
    if not path.is_file():
        raise WorkbookError("The workbook could not be found. Check the file path and try again.")
    try:
        with ZipFile(path) as archive:
            if (
                sum(i.file_size for i in archive.infolist())
                > config.MAX_EXPANDED_WORKBOOK_MB * 1024**2
            ):
                raise WorkbookError(
                    "The expanded workbook is too large. Please provide a smaller weekly workbook."
                )
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except WorkbookError:
        raise
    except (BadZipFile, OSError, ValueError, KeyError) as exc:
        raise WorkbookError(
            "The file is not a readable .xlsx workbook. If it is encrypted or damaged, save an unencrypted copy in Excel and try again."
        ) from exc
    data = ReportData(source_file=path.name)
    try:
        sheets = []
        for ws in wb:
            if ws.max_row * ws.max_column > config.MAX_WORKSHEET_CELLS:
                raise WorkbookError(
                    f"Sheet '{ws.title}' has an unusually large used range. Remove unused formatted rows/columns and save again."
                )
            sheets.append(Sheet(ws.title, list(ws.iter_rows())))
    finally:
        wb.close()
    # Formula evaluation is Excel's responsibility. Missing caches must never silently become zeros.
    formulas = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        cache_by_name = {s.title: s for s in sheets}
        missing = []
        for ws in formulas:
            for r, row in enumerate(ws.iter_rows()):
                for c, cell in enumerate(row):
                    if cell.data_type == "f" and cache_by_name[ws.title].rows[r][c].value is None:
                        missing.append(f"{ws.title}!{cell.coordinate}")
        if missing:
            data.warn(
                f"{len(missing)} formula cells have no saved result (including {', '.join(missing[:5])}). Recalculate and save in Excel/LibreOffice. Detail-based metrics are still calculated locally; missing formula values remain unavailable."
            )
    finally:
        formulas.close()
    parse_metadata(sheets, data)
    executive = find_sheet(sheets, ["Executive Summary", "Project Summary"], {"project", "overall"})
    if executive is None:
        executive = find_sheet(sheets, ["Executive Summary"], {"project", "total"})
    parse_executive_summary(executive, data)
    summary_projects = {key(p.name) for p in data.projects}
    register_sheets = [s for s in sheets if find_header(s, {"title", "status", "phase"})]
    excluded = [s.title for s in register_sheets]
    task_sheet = find_sheet(
        sheets,
        ["Detailed Task Log", "Task Log", "Weekly Tasks"],
        {"project", "title", "status"},
        excluded,
    )
    data.tasks = parse_task_log(task_sheet, data)
    breakdown = next((s for s in sheets if "projectbreakdown" in key(s.title)), None)
    if breakdown is None:
        breakdown = next(
            (
                s
                for s in sheets
                if s != task_sheet
                and s not in register_sheets
                and find_header(s, {"task_id", "title", "status"})
                and any(
                    "tasks completed this week" in clean_text(c.value).casefold()
                    for row in s.rows
                    for c in row
                )
            ),
            None,
        )
    # Parse the redundant breakdown in a separate model so it cannot introduce ghost projects.
    comparison = ReportData(
        report_start_date=data.report_start_date,
        report_end_date=data.report_end_date,
        projects=list(data.projects),
    )
    breakdown_tasks = parse_project_breakdown(breakdown, comparison)
    if not data.task_source_available and breakdown_tasks:
        data.tasks = breakdown_tasks
        data.projects = comparison.projects
        data.task_source_available = True
        data.sources["weekly_tasks"] = breakdown.title
        data.warnings.extend(w for w in comparison.warnings if w not in data.warnings)
        data.warn(
            "Detailed Task Log unavailable or headers not recognized. Weekly metrics use Project Breakdown."
        )
    elif breakdown_tasks:
        data.sources["project_breakdown"] = breakdown.title
        a = {key(t.task_id): (key(t.project), t.status, t.overdue) for t in data.tasks if t.task_id}
        b = {
            key(t.task_id): (key(t.project), t.status, t.overdue)
            for t in breakdown_tasks
            if t.task_id
        }
        if a != b:
            data.warn(
                "Project Breakdown disagrees with Detailed Task Log on task IDs, projects, statuses or overdue flags. Detailed Task Log takes precedence."
            )
    # A pipeline value is not an invoice merely because it has an Amount and Status.
    pipeline_only = [
        s.title
        for s in sheets
        if find_header(s, {"client", "opportunity"}) and not find_header(s, {"invoice_id"})
    ]
    invoices = find_sheet(
        sheets, ["Invoice Summary", "Invoices"], {"client", "amount", "status"}, pipeline_only
    )
    parse_invoice_summary(invoices, data)
    for register_sheet in register_sheets:
        parse_task_register(register_sheet, data)
    # Discard task-only project names introduced by superseded duplicate task rows.
    used_projects = (
        summary_projects
        | {key(t.project) for t in data.tasks}
        | {key(t.project) for t in data.register}
    )
    data.projects = [p for p in data.projects if key(p.name) in used_projects]
    data.sales_records = parse_sales_sheets(sheets, data)
    if not data.task_source_available and not data.projects:
        raise WorkbookError(
            'Could not find "Detailed Task Log" or another usable project/task data source. Include Project, Task Title and Status columns, or an Executive Summary with Project and Completion % / Total Tasks.'
        )
    if not data.task_source_available:
        data.warn(
            "No usable weekly task log. Only explicit Executive Summary counts are available; team workload and full task details cannot be calculated."
        )
    # Canonicalize all people once across the complete task list.
    names = {}
    for task in data.tasks:
        task.assignee = names.setdefault(key(task.assignee), task.assignee)
    unassigned = sum(t.assignee == "Unassigned" for t in data.tasks)
    if unassigned:
        data.warn(f"{unassigned} task(s) have no named assignee; grouped as Unassigned.")
    for project in data.projects:
        project_tasks = [t for t in data.tasks if t.project == project.name]
        if project_tasks and project.overall_completion is None:
            data.warn(
                f"{project.name} has weekly tasks but no overall completion percentage. Weekly completion is a separate metric."
            )
        if project.stale or any(t.stale for t in project_tasks):
            project.stale = True
            data.warn(
                f"{project.name}: data confirmation needed; task records include carried-forward status, including any recorded overdue flags."
            )
    return data
